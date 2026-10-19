import os
import openpyxl
from PIL import Image, ImageDraw, ImageFont

XLSX_PATH = "Paper/research/kimi_sources/fbi/extracted_2025/CIUS_Expanded_Homicide_Data_Table_8_Murder_Victims_by_Weapon_2021-2025.xlsx"
OUT_DIR = "Architecting_the_operation/video/yt02/archive/fbi_table8_2025"
OUT_PATH = os.path.join(OUT_DIR, "fbi_table8_2025_p1.png")

os.makedirs(OUT_DIR, exist_ok=True)
wb = openpyxl.load_workbook(XLSX_PATH)
ws = wb.active


def get_font(size, bold=False):
  candidates = (
      [
          "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
          "/System/Library/Fonts/Helvetica.ttc",
          "/Library/Fonts/Arial Bold.ttf",
      ]
      if bold
      else [
          "/System/Library/Fonts/Supplemental/Arial.ttf",
          "/System/Library/Fonts/Helvetica.ttc",
          "/Library/Fonts/Arial.ttf",
      ]
  )
  for c in candidates:
    if os.path.exists(c):
      try:
        return ImageFont.truetype(c, size)
      except Exception:
        pass
  return ImageFont.load_default()


font_title1 = get_font(34, bold=True)
font_title2 = get_font(30, bold=True)
font_title3 = get_font(26, bold=False)
font_header = get_font(24, bold=True)
font_bold = get_font(23, bold=True)
font_reg = get_font(23, bold=False)
font_sup = get_font(15, bold=False)
font_note = get_font(19, bold=False)

W, H = 2200, 1620
img = Image.new("RGB", (W, H), (255, 255, 255))
draw = ImageDraw.Draw(img)

margin_x = 90
table_w = W - 2 * margin_x
col_x = [
    margin_x,
    margin_x + 820,
    margin_x + 1060,
    margin_x + 1300,
    margin_x + 1540,
    margin_x + 1780,
    margin_x + table_w,
]

y = 65
# Title rows (rows 1-3)
t1 = str(ws.cell(1, 1).value)
t2 = str(ws.cell(2, 1).value)
t3 = str(ws.cell(3, 1).value)

draw.text((margin_x, y), t1, fill=(15, 25, 45), font=font_title1)
y += 46
draw.text((margin_x, y), t2, fill=(15, 25, 45), font=font_title2)
y += 40
draw.text((margin_x, y), t3, fill=(40, 50, 70), font=font_title3)
y += 48

# Top thick rule
draw.line([(margin_x, y), (margin_x + table_w, y)], fill=(20, 30, 50), width=4)
y += 4

# Header row (row 4)
row_h = 52
draw.rectangle(
    [margin_x, y, margin_x + table_w, y + row_h], fill=(238, 242, 248)
)
headers = [ws.cell(4, c).value for c in range(1, 7)]
draw.text(
    (col_x[0] + 16, y + 13), str(headers[0]), fill=(15, 25, 45), font=font_header
)
for i in range(1, 6):
  txt = str(headers[i])
  tw = draw.textlength(txt, font=font_header)
  draw.text(
      (col_x[i + 1] - 20 - tw, y + 13), txt, fill=(15, 25, 45), font=font_header
  )
y += row_h
draw.line([(margin_x, y), (margin_x + table_w, y)], fill=(20, 30, 50), width=3)

# Data rows (rows 5 to 22)
firearm_subrows = {7, 8, 9, 10, 11}
bboxes = {}

for r in range(5, 23):
  label = str(ws.cell(r, 1).value)
  vals = [ws.cell(r, c).value for c in range(2, 7)]
  rh = 48
  # Zebra striping
  if r == 8:
    # Subtle warm tint on Rifles row or clean white/stripe
    bg = (255, 249, 230)
  elif r % 2 == 1:
    bg = (249, 250, 252)
  else:
    bg = (255, 255, 255)
  draw.rectangle([margin_x, y, margin_x + table_w, y + rh], fill=bg)

  is_bold_row = r in (5, 6)
  f_use = font_bold if is_bold_row else font_reg
  indent = 44 if r in firearm_subrows else 16

  # Handle superscripts in label cleanly
  if label.endswith("1,2"):
    base_lbl = label[:-3]
    sup_lbl = "1,2"
  elif label.endswith("2") and r in (19, 20):
    base_lbl = label[:-1]
    sup_lbl = "2"
  else:
    base_lbl = label
    sup_lbl = ""

  draw.text((col_x[0] + indent, y + 11), base_lbl, fill=(20, 25, 35), font=f_use)
  if sup_lbl:
    bw = draw.textlength(base_lbl, font=f_use)
    draw.text(
        (col_x[0] + indent + bw + 2, y + 8),
        sup_lbl,
        fill=(20, 25, 35),
        font=font_sup,
    )

  for i, v in enumerate(vals):
    v_str = f"{int(v):,}" if isinstance(v, (int, float)) else str(v)
    tw = draw.textlength(v_str, font=f_use)
    draw.text(
        (col_x[i + 2] - 20 - tw, y + 11), v_str, fill=(20, 25, 35), font=f_use
    )

  bboxes[base_lbl] = (margin_x, y, table_w, rh)
  y += rh
  line_w = 2 if r in (5, 11, 22) else 1
  line_col = (40, 50, 70) if r == 22 else (205, 212, 222)
  draw.line(
      [(margin_x, y), (margin_x + table_w, y)], fill=line_col, width=line_w
  )

# Bottom border
draw.line([(margin_x, y), (margin_x + table_w, y)], fill=(20, 30, 50), width=3)
y += 22

# Footnotes (rows 23-25)
for r in (23, 24, 25):
  fn = str(ws.cell(r, 1).value)
  draw.text((margin_x, y), fn, fill=(55, 65, 80), font=font_note)
  y += 32

y += 10
draw.text(
    (margin_x, y),
    "Source: Federal Bureau of Investigation, Crime in the United States /"
    " Crime Data Explorer (CIUS Expanded Homicide Data Table 8, 2021–2025).",
    fill=(90, 100, 115),
    font=font_note,
)

img.save(OUT_PATH, format="PNG")
print("Saved:", OUT_PATH, img.size)
for k in ["Total", "Total firearms:", "Handguns", "Rifles", "Shotguns"]:
  print(f"BBox {k}: {bboxes.get(k)}")
